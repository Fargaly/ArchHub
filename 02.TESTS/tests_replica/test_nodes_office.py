"""Synthetic host-boundary courts: no real Outlook, Excel UI or client files."""
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from nodelang.pipeline_engines import PIPELINE_ENGINES
from nodelang.library_engines import library_catalogue


ENGINES = (
    'outlook.search', 'outlook.read_thread', 'outlook.reply_draft',
    'outlook.forward_draft', 'library.draft_email', 'excel.read_table',
    'excel.build_workbook', 'excel.write_cells', 'excel.share_workbook',
    'excel.open', 'pdf.extract_text', 'pdf.snip',
)


@pytest.mark.parametrize('engine', ENGINES)
def test_engine_is_visible_with_typed_editable_sockets(engine):
    assert engine in PIPELINE_ENGINES, 'engine missing: ' + engine
    rows = [r for g in library_catalogue() for r in g['items']
            if r['engine'] == engine and not r.get('hidden')]
    assert rows, 'visible library engine missing: ' + engine
    row = rows[0]
    assert row['params'], 'parameters missing: ' + engine
    targets = {p['id']: p for p in row['ports']['target']}
    for name in row['params']:
        assert name in targets, 'parameter socket missing: ' + name
        assert targets[name]['type'] not in ('any', 'unknown')
    assert all(p['type'] not in ('any', 'unknown')
               for side in row['ports'].values() for p in side)


class Collection:
    def __init__(self, rows=()):
        self.rows = list(rows)

    @property
    def Count(self):
        return len(self.rows)

    def Item(self, index):
        return self.rows[index - 1]

    def Sort(self, *args):
        pass

    def Add(self, path):
        item = SimpleNamespace(FileName=Path(path).name, Size=Path(path).stat().st_size,
                               PropertyAccessor=SimpleNamespace(SetProperty=lambda *args: None))
        self.rows.append(item)
        return item


class Mail:
    def __init__(self, id='message-1', subject='coordination', body='evidence', conversation='thread-1'):
        self.EntryID, self.Subject, self.Body = id, subject, body
        self.ConversationID, self.SenderName = conversation, 'Synthetic sender'
        self.SenderEmailAddress, self.ReceivedTime = 'sender@example.invalid', '2026-10-25T09:00:00+00:00'
        self.UnRead, self.Class = True, 43
        self.To, self.CC, self.HTMLBody = 'team@example.invalid', '', '<html><body>QUOTED</body></html>'
        self.Attachments, self.saved, self.displayed = Collection(), False, False
        self.Parent = SimpleNamespace(StoreID='store-1', EntryID='inbox')

    def ReplyAll(self):
        self.reply = Mail('reply')
        return self.reply

    def Forward(self):
        self.forward = Mail('forward')
        return self.forward

    def Save(self):
        self.saved = True

    def Display(self):
        self.displayed = True

    def Send(self):
        pytest.fail('send path called')


@pytest.fixture
def outlook(monkeypatch):
    from nodelang import library_engines as library
    from nodelang import host_brokers as hosts
    mail = Mail()
    other = Mail('other', conversation='thread-2')
    folder = SimpleNamespace(EntryID='inbox', StoreID='store-1', Items=Collection([mail, other]), Name='Inbox')
    ns = SimpleNamespace(GetDefaultFolder=lambda _: folder,
                         GetFolderFromID=lambda *args: folder,
                         GetItemFromID=lambda *args: mail)
    created = []
    def create(_):
        item = Mail('draft')
        created.append(item)
        return item
    app = SimpleNamespace(GetNamespace=lambda _: ns, CreateItem=create)
    monkeypatch.setattr(library, '_OUTLOOK', [lambda: app])
    monkeypatch.setattr(hosts, '_outlook_application', lambda: app)
    monkeypatch.setattr(hosts, '_com_alive', lambda _: True)
    return mail, created


def run(engine, params=None, feeds=None):
    params = dict(params or {})
    if 'grant' in params and params['grant'].get('scope') == 'tree':
        from nodelang.clean_office_adapter import coordination_arguments, coordination_write_request
        from nodelang.nodeslib_v2_engines import _digest_payload
        args = coordination_arguments(engine, params, feeds or {})
        destination = args.get('output') if engine == 'pdf.snip' else args.get('destination') if engine == 'excel.share_workbook' else args['path']
        assert Path(params['grant']['path']).resolve() in Path(destination).resolve().parents
        params['grant'] = {**params['grant'], 'path': str(destination), 'scope': 'exact',
                           'input_digest': _digest_payload(coordination_write_request(engine, args))}
    return PIPELINE_ENGINES[engine](params, feeds or {})[0]


def test_search_and_thread_preserve_unread_and_bound_body_reads(outlook):
    mail, _ = outlook
    result = run('outlook.search', {'folders': ['inbox'], 'since': '2026-10-24',
                 'keywords': ['coordination'], 'include_body': False, 'include_attachments': True})
    assert len(result['out']) == 2
    assert 'body' not in result['out'][0]
    assert result['out'][0]['id'] == mail.EntryID
    thread = run('outlook.read_thread', {'conversation_id': 'thread-1', 'folders': ['inbox']})
    assert [r['id'] for r in thread['out']] == [mail.EntryID]
    assert mail.UnRead is True


@pytest.mark.parametrize('engine,kind', [('library.draft_email', 'new'),
                         ('outlook.reply_draft', 'reply'), ('outlook.forward_draft', 'forward')])
def test_drafts_keep_signature_caps_url_attachments_and_never_send(engine, kind, outlook, tmp_path):
    mail, created = outlook
    signature = tmp_path / 'HOUSE.htm'
    signature.write_text('<html><body><div>{{body}}</div><b>HOUSE SIGNATURE</b></body></html>')
    attachment = tmp_path / 'proof.txt'
    attachment.write_text('synthetic')
    params = {'to': 'team@example.invalid', 'cc': 'cc@example.invalid', 'subject': 'review required',
              'body': '<p>See <a href="https://example.invalid/proof">click here</a></p>',
              'body_format': 'html', 'signature': 'HOUSE', 'signature_dir': str(tmp_path),
              'attachments': [str(attachment)], 'message_id': mail.EntryID, 'store_id': 'store-1'}
    answer = run(engine, params)
    draft = created[0] if kind == 'new' else getattr(mail, kind)
    assert answer['out']['draft_only'] is True
    assert draft.saved and not draft.displayed
    assert draft.Subject == 'REVIEW REQUIRED'
    assert '<b>HOUSE SIGNATURE</b>' in draft.HTMLBody
    assert '>https://example.invalid/proof</a>' in draft.HTMLBody
    assert '{{body}}' not in draft.HTMLBody
    assert draft.CC == 'cc@example.invalid' and draft.Attachments.Count == 1
    if kind != 'new':
        assert 'QUOTED' in draft.HTMLBody
    with pytest.raises(ValueError, match='draft.only|send'):
        run(engine, {**params, 'send': True})


@pytest.fixture
def file_authority(monkeypatch, tmp_path):
    from nodelang import nodeslib_v2_engines as policy
    monkeypatch.setattr(policy, '_WRITE_GRANT_VERIFIER', lambda request: {'ok': True})
    import time
    return {'kind': 'cde-write-permit', 'path': str(tmp_path), 'scope': 'tree',
            'operations': ['write_file'], 'expires_at': time.time() + 300,
            'digest': 'a' * 64, 'caller': 'court', 'session': 'court', 'work': 'court',
            'action': 'write_file'}


def test_excel_build_read_write_share_and_images(tmp_path, file_authority, monkeypatch):
    import openpyxl
    from PIL import Image
    from nodelang import clean_office_adapter as office
    monkeypatch.setattr(office, '_running', lambda *args: (_ for _ in ()).throw(office.OfficeUnreachable('no Excel')))
    picture = tmp_path / 'proof.png'
    Image.new('RGB', (40, 20), 'red').save(picture)
    path = tmp_path / 'tracker.xlsx'
    spec = [{'name': 'Comments', 'columns': ['ID', 'Comment', 'Status', 'Evidence'],
             'rows': [{'ID': 'C-1', 'Comment': 'Ramp checked', 'Status': 'Open'}],
             'validation': {'C2:C20': ['Open', 'Closed']},
             'formulas': {'F2': '=COUNTA(A2:A20)'},
             'images': [{'cell': 'D2', 'path': str(picture)}]}]
    params = {'path': str(path), 'sheets': spec, 'grant': file_authority}
    run('excel.build_workbook', params)
    read = run('excel.read_table', {'path': str(path), 'id_column_regex': '^ID$', 'id_regex': '^C-'})
    assert read['rows'][0]['ID'] == 'C-1' and read['tables'][0]['text_column'] == 'Comment'
    run('excel.write_cells', {'path': str(path), 'cells': [
        {'sheet': 'Comments', 'cell': 'C2', 'text': 'Closed', 'fill': 'FFFF00'},
        {'sheet': 'Comments', 'cell': 'D3', 'image': str(picture)}], 'grant': file_authority})
    book = openpyxl.load_workbook(path)
    sheet = book['Comments']
    assert sheet['C2'].value == 'Closed' and sheet['C2'].fill.fgColor.rgb.endswith('FFFF00')
    assert len(sheet._images) == 2 and sheet.auto_filter.ref
    assert sheet['B2'].alignment.wrap_text and len(sheet.data_validations.dataValidation) == 1
    assert sheet['F2'].value == '=COUNTA(A2:A20)'
    book.close()
    shared = tmp_path / 'SHARED' / path.name
    run('excel.share_workbook', {'path': str(path), 'destination': str(shared), 'grant': file_authority})
    assert shared.read_bytes() == path.read_bytes()
    assert run('excel.open', {'path': str(path)})['out']['opened'] is False


@pytest.fixture
def delivered_pdf(tmp_path):
    # A real one-page PDF fixture; no second PDF library is needed to create it.
    stream = b'BT /F1 12 Tf 72 720 Td (Ramp checked in delivered package) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
               b'/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    pdf, offsets = bytearray(b'%PDF-1.4\n'), [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(str(number).encode() + b' 0 obj\n' + obj + b'\nendobj\n')
    start = len(pdf)
    pdf.extend(b'xref\n0 6\n0000000000 65535 f \n')
    for offset in offsets[1:]:
        pdf.extend(('%010d 00000 n \n' % offset).encode())
    pdf.extend(b'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n' + str(start).encode() + b'\n%%EOF\n')
    path = tmp_path / 'delivered.pdf'
    path.write_bytes(pdf)
    return path


def test_pdf_text_and_phrase_snip(tmp_path, file_authority, delivered_pdf):
    from PIL import Image
    path = delivered_pdf
    text = run('pdf.extract_text', {'path': str(path)})
    assert text['out'][0]['page'] == 1 and 'Ramp checked' in text['out'][0]['text']
    snip = tmp_path / 'snip.png'
    out = run('pdf.snip', {'path': str(path), 'page': 1, 'phrase': 'Ramp checked',
                         'output': str(snip), 'grant': file_authority})
    assert out['file'] == str(snip) and out['out']['page'] == 1
    with Image.open(snip) as image:
        assert image.width > 10 and image.height > 10
    with pytest.raises(ValueError, match='not found'):
        run('pdf.snip', {'path': str(path), 'phrase': 'missing', 'output': str(snip), 'grant': file_authority})


def test_d1_runs_only_registered_nodes_with_real_xlsx_pdf_images_and_saved_draft(tmp_path, file_authority, outlook, monkeypatch, delivered_pdf):
    import openpyxl
    from nodelang import clean_office_adapter as office
    from nodelang.nodeslib_v2_engines import _digest_payload
    from nodelang.stem_graph_evaluation import StemNode, StemWire
    from nodelang.map_import import resolve_map_path
    from nodelang.universal_application import build_universal_application, project_universal_canvas, connect_universal_roots
    from nodelang.universal_pipeline import create_engine_node, run_universal_pipeline
    monkeypatch.setattr(office, '_running', lambda *args: (_ for _ in ()).throw(office.OfficeUnreachable('no Excel')))
    comments, delivered = tmp_path / 'comments.xlsx', delivered_pdf
    snip, tracker = tmp_path / 'evidence.png', tmp_path / 'tracker.xlsx'
    book = openpyxl.Workbook()
    book.active.append(['ID', 'Comment', 'Status'])
    book.active.append(['C-1', 'Ramp checked', 'Open'])
    book.active.append(['C-2', 'Unrelated', 'Closed'])
    book.save(comments)
    book.close()
    (tmp_path / 'HOUSE.htm').write_text('<html><body>{{body}}<b>HOUSE SIGNATURE</b></body></html>')
    snip_params = {'path': str(delivered), 'output': str(snip)}
    build_params = {'path': str(tracker), 'columns': ['ID', 'Comment', 'Status', 'Evidence']}
    expected_rows = [{'ID': 'C-1', 'Comment': 'Ramp checked', 'Status': 'Open'}]
    for engine, params, dynamic, destination in (
        ('pdf.snip', snip_params, {'phrase': 'Ramp checked'}, snip),
        ('excel.build_workbook', build_params, {'rows': expected_rows, 'images': [{'path': str(snip), 'cell': 'D2'}]}, tracker),
    ):
        args = office.coordination_arguments(engine, params, dynamic)
        params['grant'] = {**file_authority, 'scope': 'exact', 'path': str(destination),
                           'input_digest': _digest_payload(office.coordination_write_request(engine, args))}
    nodes = [
        StemNode('comments', 'excel.read_table', {'path': str(comments)}),
        StemNode('filter', 'shape.filter', {'predicate': "item.Status == 'Open'"}),
        StemNode('comment', 'shape.pluck', {'field': 'Comment'}),
        StemNode('phrase', 'text.ops', {'operation': 'join', 'separator': ''}),
        StemNode('delivered', 'pdf.extract_text', {'path': str(delivered)}),
        StemNode('mail_history', 'outlook.search', {'keywords': ['coordination']}),
        StemNode('evidence', 'pdf.snip', snip_params),
        StemNode('tracker', 'excel.build_workbook', build_params),
        StemNode('draft', 'library.draft_email', {'to': 'team@example.invalid', 'subject': 'evidenced responses',
                 'body': '<a href="https://example.invalid/delivery">delivery</a>', 'body_format': 'html',
                 'signature': 'HOUSE', 'signature_dir': str(tmp_path)}),
    ]
    wires = [StemWire(*wire) for wire in (
        ('comments', 'rows', 'filter', 'items'), ('filter', 'items_out', 'comment', 'items'),
        ('comment', 'values', 'phrase', 'items'), ('phrase', 'text', 'evidence', 'phrase'),
        ('filter', 'items_out', 'tracker', 'rows'), ('evidence', 'images', 'tracker', 'images'),
        ('tracker', 'attachments', 'draft', 'attachments'),
    )]
    store, registry = build_universal_application(resolve_map_path())
    try:
        roots = {}
        for node in nodes:
            properties = {key: json.dumps(value) if isinstance(value, (dict, list, bool)) else value
                          for key, value in node.parameters.items()}
            roots[node.root_id] = create_engine_node(store, registry, title=node.root_id,
                engine=node.engine, properties=properties)['root']
        projected = {node['id']: node for node in project_universal_canvas(store, registry)['nodes']}
        for wire in wires:
            source, target = roots[wire.source], roots[wire.target]
            output = next(p['id'] for p in projected[source]['ports'] if p.get('name') == wire.source_interface and p['side'] == 'source')
            input_port = next(p['id'] for p in projected[target]['ports'] if p.get('name') == wire.target_interface and p['side'] == 'target')
            connect_universal_roots(store, registry, source, target, source_interface=output, target_interface=input_port)
        answer = run_universal_pipeline(store, registry, effect_engines=PIPELINE_ENGINES, only_roots=list(roots.values()))
        result = SimpleNamespace(pending=answer['pending'], node_outputs={name: answer['outputs'].get(root, {}) for name, root in roots.items()})
        assert answer['ran'] == len(nodes) and answer['wires'] == len(wires)
    finally:
        store.close()
    assert not result.pending, result.pending
    assert len(result.node_outputs['mail_history']['out']) == 2
    assert 'Ramp checked' in result.node_outputs['delivered']['out'][0]['text']
    assert result.node_outputs['draft']['out']['draft_only'] is True
    output = openpyxl.load_workbook(tracker)
    assert output.active['A2'].value == 'C-1' and output.active.max_row == 2
    assert len(output.active._images) == 1
    output.close()
    draft = outlook[1][0]
    assert draft.saved and draft.Subject == 'EVIDENCED RESPONSES' and draft.Attachments.Count == 1
    assert 'HOUSE SIGNATURE' in draft.HTMLBody and '>https://example.invalid/delivery</a>' in draft.HTMLBody
    print('D1 CANVAS: %d persisted nodes / %d wires / saved draft with image workbook' % (len(nodes), len(wires)))


def test_open_shared_workbook_uses_exact_existing_com_book_and_reads_back(tmp_path, file_authority, monkeypatch):
    from nodelang import clean_office_adapter as office
    path = tmp_path / 'shared.xlsx'
    cell = SimpleNamespace(Value2='', NumberFormat='', Interior=SimpleNamespace(Color=0))
    sheet = SimpleNamespace(Range=lambda _: cell)
    saved = []
    book = SimpleNamespace(FullName=str(path), Worksheets=SimpleNamespace(Item=lambda _: sheet), Save=lambda: saved.append(True))
    app = SimpleNamespace(Workbooks=Collection([SimpleNamespace(FullName=str(tmp_path / 'other.xlsx')), book]))
    monkeypatch.setattr(office, '_running', lambda *args: app)
    run('excel.write_cells', {'path': str(path), 'shared': True, 'cells': [
        {'cell': 'B2', 'text': '=literal', 'fill': 'FF0000'}], 'grant': file_authority})
    assert cell.Value2 == '=literal' and cell.NumberFormat == '@'
    assert cell.Interior.Color == 255 and saved == [True]


def test_locked_workbook_retries_without_replacing_original(tmp_path, file_authority, monkeypatch):
    import openpyxl
    from nodelang import clean_office_adapter as office
    path = tmp_path / 'locked.xlsx'
    book = openpyxl.Workbook()
    book.save(path)
    book.close()
    original = path.read_bytes()
    monkeypatch.setattr(office, '_excel_open_book', lambda _: None)
    calls = []
    def locked(*args, **kwargs):
        calls.append(1)
        raise PermissionError('synthetic lock')
    monkeypatch.setattr(office, '_atomic_office_output', locked)
    with pytest.raises(PermissionError, match='synthetic lock'):
        run('excel.write_cells', {'path': str(path), 'cells': [{'cell': 'A1', 'text': 'new'}],
             'retry_count': 2, 'retry_seconds': 0, 'grant': file_authority})
    assert len(calls) == 3 and path.read_bytes() == original


def test_write_authority_rejects_unbound_grant_before_creating_any_file(tmp_path):
    path = tmp_path / 'unauthorized.xlsx'
    from nodelang import workspace_roots_catalogue as roots
    from unittest.mock import patch
    with patch.object(roots, 'read_registered_roots', return_value=({}, {})), pytest.raises(ValueError, match='owner Canvas run'):
        run('excel.build_workbook', {'path': str(path), 'columns': ['ID']})
    assert not path.exists()


def test_canvas_library_placement_keeps_parameters_sockets_wires_and_undo_redo():
    from nodelang.map_import import resolve_map_path
    from nodelang.universal_application import (build_universal_application, project_universal_canvas,
        connect_universal_roots, undo_universal_change, redo_universal_change)
    from nodelang.universal_pipeline import create_engine_node, read_engine_library, seed_engine_library
    from nodelang.clean_office_adapter import COORDINATION_NODES, coordination_arguments
    store, registry = build_universal_application(resolve_map_path())
    try:
        seed_engine_library(store, registry)
        visible = {row['engine']: row for group in read_engine_library(store.snapshot())
                   for row in group['items'] if not row.get('hidden')}
        roots = {}
        for engine in ENGINES:
            entry = visible[engine]
            root = create_engine_node(store, registry, title=entry['title'], engine=engine, item=entry['id'])['root']
            roots[engine] = root
            projection = project_universal_canvas(store, registry)
            node = next(n for n in projection['nodes'] if n['id'] == root)
            properties = {r['label']: r['value'] for r in node['params']}
            assert set(COORDINATION_NODES[engine]['params']) <= set(properties)
            if 'path' in properties:
                properties['path'] = 'synthetic.xlsx'
            coordination_arguments(engine, properties, {})
            ports = {p.get('name') for p in node['ports'] if p.get('mode') == 'connection' and p['side'] == 'target'}
            assert set(COORDINATION_NODES[engine]['params']) <= ports
        projection = project_universal_canvas(store, registry)
        source = next(n for n in projection['nodes'] if n['id'] == roots['pdf.snip'])
        target = next(n for n in projection['nodes'] if n['id'] == roots['excel.build_workbook'])
        source_port = next(p['id'] for p in source['ports'] if p.get('name') == 'images' and p['side'] == 'source')
        target_port = next(p['id'] for p in target['ports'] if p.get('name') == 'images' and p['side'] == 'target')
        wire, _ = connect_universal_roots(store, registry, source['id'], target['id'],
                                         source_interface=source_port, target_interface=target_port)
        assert wire in {w['id'] for w in project_universal_canvas(store, registry)['wires']}
        undo_universal_change(store, registry)
        assert wire not in {w['id'] for w in project_universal_canvas(store, registry)['wires']}
        redo_universal_change(store, registry)
        assert wire in {w['id'] for w in project_universal_canvas(store, registry)['wires']}
        print('CANVAS: 12 visible placements; every parameter socket; evidence image wire; Undo/Redo')
    finally:
        store.close()


def test_selected_store_uses_that_stores_default_folders(outlook, monkeypatch):
    from nodelang import library_engines as library
    app = library._OUTLOOK[0]()
    namespace = app.GetNamespace('MAPI')
    selected = []
    folder = namespace.GetDefaultFolder(6)
    namespace.GetStoreFromID = lambda store: SimpleNamespace(GetDefaultFolder=lambda number: selected.append((store, number)) or folder)
    namespace.GetFolderFromID = lambda *args: pytest.fail('folder alias treated as EntryID')
    answer = run('outlook.search', {'store_id': 'selected-store'})
    assert answer['out'] and selected == [('selected-store', 6), ('selected-store', 5)]


def test_signature_fragment_keeps_quote_and_embeds_relative_logo(outlook, tmp_path):
    from PIL import Image
    assets = tmp_path / 'HOUSE_FILES'
    assets.mkdir()
    Image.new('RGB', (8, 8)).save(assets / 'logo.png')
    (tmp_path / 'HOUSE.htm').write_bytes(b'<meta charset="windows-1252"><b>HOUSE SIGNATURE \xe9</b><img src="HOUSE_FILES/logo.png">')
    run('outlook.reply_draft', {'message_id': 'message-1', 'body': 'New content',
        'signature': 'HOUSE', 'signature_dir': str(tmp_path)})
    draft = outlook[0].reply
    assert 'QUOTED' in draft.HTMLBody and 'HOUSE SIGNATURE \u00e9' in draft.HTMLBody
    assert 'cid:archhub-signature-0@local' in draft.HTMLBody and draft.Attachments.Count == 1


def test_open_never_launches_an_executable(tmp_path, monkeypatch):
    from nodelang import clean_office_adapter as office
    file = tmp_path / 'program.exe'
    file.write_bytes(b'not a real executable')
    monkeypatch.setattr(office.os, 'startfile', lambda *args: pytest.fail('executable launched'))
    with pytest.raises(ValueError, match='workbook'):
        run('excel.open', {'path': str(file), 'show': True})


def test_unsupported_drawings_refuse_without_losing_any_package_bytes(tmp_path, file_authority, monkeypatch):
    import openpyxl
    import zipfile
    from nodelang import clean_office_adapter as office
    path = tmp_path / 'shapes.xlsx'
    book = openpyxl.Workbook()
    book.save(path)
    book.close()
    with zipfile.ZipFile(path, 'a') as archive:
        archive.writestr('xl/drawings/preserved.xml', '<drawing><sp>untouched synthetic shape</sp></drawing>')
    before = path.read_bytes()
    monkeypatch.setattr(office, '_excel_open_book', lambda _: None)
    with pytest.raises(office.OfficeUnreachable, match='preserve'):
        run('excel.write_cells', {'path': str(path), 'cells': [{'cell': 'A1', 'text': 'change'}], 'grant': file_authority})
    assert path.read_bytes() == before


def test_clean_canvas_uses_existing_owner_gate_and_private_workspace_without_fake_grant(tmp_path, monkeypatch):
    from nodelang.application_server import _CleanAuthorityHttpServer
    from nodelang import workspace_roots_catalogue as roots
    from nodelang import clean_office_adapter as office
    calls = []
    monkeypatch.setattr(roots, 'read_registered_roots', lambda: ({'SYNTHETIC': {'path': str(tmp_path)}}, {}))
    monkeypatch.setattr(roots, 'root_bound_registration', lambda path: ('workspace-root:SYNTHETIC', 'digest', {'privacy': 'private', 'path': str(tmp_path)}))
    def gate(name, effect):
        def admitted(params, feeds):
            calls.append(name)
            return effect(params, feeds)
        return admitted
    table = _CleanAuthorityHttpServer._clean_engine_table(gate)
    path = tmp_path / 'owner.xlsx'
    result, _ = table['excel.build_workbook']({'path': str(path), 'columns': ['ID'], 'rows': [['C-1']]}, {})
    assert path.is_file() and result['file'] == str(path) and calls == ['excel.build_workbook']
    refused = _CleanAuthorityHttpServer._clean_engine_table()
    with pytest.raises(ValueError, match='no adapter'):
        refused['excel.build_workbook']({'path': str(tmp_path / 'refused.xlsx')}, {})
    assert not (tmp_path / 'refused.xlsx').exists()
    def denied_gate(_name, _effect):
        def denied(*args):
            raise ValueError('revoked or non-owner')
        return denied
    with pytest.raises(ValueError, match='revoked or non-owner'):
        _CleanAuthorityHttpServer._clean_engine_table(denied_gate)['excel.build_workbook']({'path': str(path)}, {})
    monkeypatch.setattr(roots, 'root_bound_registration', lambda path: ('workspace-root:SYNTHETIC', 'digest', {'privacy': 'public'}))
    with pytest.raises(ValueError, match='private registered'):
        table['excel.build_workbook']({'path': str(tmp_path / 'public.xlsx'), 'columns': ['ID']}, {})
    with pytest.raises(ValueError, match='owner Canvas run'):
        office._authorize_coordination_write('excel.build_workbook', {}, path)


def test_real_clean_owner_gate_refuses_changed_session_and_other_account(tmp_path, monkeypatch):
    from nodelang.application_server import _CleanAuthorityHttpServer
    from nodelang import workspace_roots_catalogue as roots
    from nodelang import clean_office_adapter as office
    binding = SimpleNamespace(session_root='session-1', subject_root='owner')
    server = SimpleNamespace(authority=SimpleNamespace(manifest=SimpleNamespace(principal_root='owner')),
                             _resolve_binding=lambda *args, **kwargs: binding)
    monkeypatch.setattr(roots, 'read_registered_roots', lambda: ({'SYNTHETIC': {'path': str(tmp_path)}}, {}))
    monkeypatch.setattr(roots, 'root_bound_registration', lambda path: ('root', 'digest', {'privacy': 'private', 'path': str(tmp_path)}))
    records = []
    gate = _CleanAuthorityHttpServer._clean_owner_effect_gate(server, 'token', 'csrf', 'session-1', records)
    table = _CleanAuthorityHttpServer._clean_engine_table(gate)
    path = tmp_path / 'allowed.xlsx'
    table['excel.build_workbook']({'path': str(path), 'columns': ['ID']}, {})
    assert path.is_file() and records[-1]['succeeded']
    binding.session_root = 'revoked-session'
    with pytest.raises(Exception, match='Browser binding changed'):
        table['excel.build_workbook']({'path': str(tmp_path / 'revoked.xlsx'), 'columns': ['ID']}, {})
    binding.session_root, binding.subject_root = 'session-1', 'other-account'
    with pytest.raises(ValueError, match='application owner'):
        table['excel.build_workbook']({'path': str(tmp_path / 'other.xlsx'), 'columns': ['ID']}, {})
    assert not (tmp_path / 'revoked.xlsx').exists() and not (tmp_path / 'other.xlsx').exists()
    monkeypatch.setattr(roots, 'root_bound_registration', lambda path: ('root', 'digest', {'privacy': 'private', 'path': str(tmp_path / 'MOVED')}))
    with pytest.raises(ValueError, match='changed before'):
        office._authorize_coordination_write('excel.build_workbook', {}, tmp_path / 'old.xlsx', owner_admitted=True)


def test_metadata_search_never_reads_body_properties(outlook, monkeypatch):
    original = Mail.__getattribute__
    def no_body(self, name):
        if name in ('Body', 'HTMLBody'):
            pytest.fail('metadata search read a message body')
        return original(self, name)
    monkeypatch.setattr(Mail, '__getattribute__', no_body)
    assert run('outlook.search', {'include_body': False})['out']


def test_default_outlook_html_signature_and_reply_subject_are_preserved(outlook, monkeypatch):
    from nodelang import library_engines as library
    app = library._OUTLOOK[0]()
    created = []
    def create(_):
        mail = Mail('draft')
        mail.GetInspector = object()
        mail.HTMLBody = '<html><body><b>DEFAULT SIGNATURE</b></body></html>'
        created.append(mail)
        return mail
    monkeypatch.setattr(app, 'CreateItem', create)
    run('library.draft_email', {'body': 'New content'})
    assert 'DEFAULT SIGNATURE' in created[0].HTMLBody and 'New content' in created[0].HTMLBody
    assert created[0].BodyFormat == 2 and not created[0].displayed
    run('outlook.reply_draft', {'message_id': 'message-1', 'body': 'Response'})
    assert outlook[0].reply.Subject == 'COORDINATION'


def test_table_dates_remain_json_serializable_for_the_canvas_route(tmp_path):
    import datetime
    import openpyxl
    path = tmp_path / 'dates.xlsx'
    book = openpyxl.Workbook()
    book.active.append(['ID', 'Comment', 'Received', 'Elapsed'])
    duration = datetime.timedelta(days=1, hours=3, minutes=15)
    book.active.append(['C-1', 'Comment', datetime.datetime(2026, 10, 25, 9, 30), duration])
    book.active['D2'].number_format = '[h]:mm:ss'
    book.save(path)
    book.close()
    answer = run('excel.read_table', {'path': str(path)})
    assert answer['rows'][0]['Received'] == '2026-10-25T09:30:00'
    assert answer['rows'][0]['Elapsed'] == str(duration)
    assert json.loads(json.dumps(answer))['rows'][0]['ID'] == 'C-1'


def test_existing_inbox_and_office_range_engines_keep_their_behavior(outlook, tmp_path):
    from nodelang.nodeslib_v2_engines import _write_xlsx
    answer = run('outlook.inbox', {'count': 1})
    assert len(answer['out']) == 1 and answer['out'][0]['subject'] == 'coordination'
    assert answer['out'][0]['unread'] is True
    path = tmp_path / 'legacy.xlsx'
    _write_xlsx([{'ID': 'C-1'}], path)
    result = run('office.read', {'operation': 'excel.read_range', 'file': str(path),
                               'sheet': 'Sheet1', 'range': 'A1:A2'})
    assert result['out']['rows'] == [['ID'], ['C-1']]


def test_draft_keeps_wired_plain_text_and_running_outlook_only(outlook, monkeypatch):
    from nodelang import library_engines as library
    run('library.draft_email', {'subject': 'wired'}, {'in': 'A & B\nNext line'})
    assert 'A &amp; B<br>Next line' in outlook[1][0].HTMLBody
    assert not outlook[1][0].displayed
    monkeypatch.setattr(library, '_OUTLOOK', [lambda: None])
    assert run('library.draft_email', {'body': 'No launch'})['out'] == []
    assert run('outlook.search')['ok'] is False


@pytest.mark.parametrize('markup,url', [
    ('<a href=https://evil.example/x>https://good.example</a>', 'https://evil.example/x'),
    ('<A HREF=HTTPS://example.invalid/x>wrong</A>', 'HTTPS://example.invalid/x'),
    ('<a title="1 > 0" HrEf="https://example.invalid/?a=1&amp;b=2"><b>wrong</b></a>',
     'https://example.invalid/?a=1&b=2'),
    ("<a href='https://example.invalid/single'>wrong</a>", 'https://example.invalid/single'),
    ('<a href=mailto:team@example.invalid>wrong</a>', 'mailto:team@example.invalid'),
    ('<p>A\rB</p>\n<a href=https://evil.example/x>wrong</a>', 'https://evil.example/x'),
])
def test_draft_link_label_matches_parsed_href(outlook, markup, url):
    from html import escape
    run('library.draft_email', {'body_format': 'html', 'body': '<p>BEFORE</p>' + markup + '<p>AFTER</p>'})
    body = outlook[1][0].HTMLBody
    assert '>' + escape(url) + '</' in body
    assert '<p>BEFORE</p>' in body and '<p>AFTER</p>' in body
    assert 'https://good.example' not in body and '>wrong<' not in body


def test_self_closing_draft_anchor_is_refused_before_outlook(monkeypatch):
    from nodelang import library_engines as library
    monkeypatch.setattr(library, '_OUTLOOK', [lambda: pytest.fail('Outlook accessed before link validation')])
    with pytest.raises(ValueError, match='self.closing'):
        run('library.draft_email', {'body_format': 'html',
            'body': '<a href="https://evil.example/x"/>wrong</a>'})


def test_workbook_formula_like_header_is_literal_text(tmp_path, file_authority):
    import openpyxl
    path = tmp_path / 'headers.xlsx'
    run('excel.build_workbook', {'path': str(path), 'columns': ['=SUM(1,2)', 'Comment'],
        'rows': [['=literal', 'evidence']], 'formulas': {'C2': '=SUM(1,2)'}, 'grant': file_authority})
    book = openpyxl.load_workbook(path)
    try:
        assert book.active['A1'].value == '=SUM(1,2)' and book.active['A1'].data_type == 's'
        assert book.active['A2'].value == '=literal' and book.active['A2'].data_type == 's'
        assert book.active['C2'].data_type == 'f'
    finally:
        book.close()


@pytest.mark.parametrize('engine', ['library.draft_email', 'outlook.reply_draft', 'outlook.forward_draft'])
@pytest.mark.parametrize('name', ['Send', 'SEND', 'sEnD'])
@pytest.mark.parametrize('source', ['params', 'feeds'])
def test_send_parameter_any_case_is_refused_before_outlook(engine, name, source, monkeypatch):
    from nodelang import library_engines as library
    monkeypatch.setattr(library, '_OUTLOOK', [lambda: pytest.fail('Outlook accessed before Send refusal')])
    params, feeds = {'body': 'draft'}, {}
    (params if source == 'params' else feeds)[name] = False
    with pytest.raises(ValueError, match='draft.only|send'):
        run(engine, params, feeds)
